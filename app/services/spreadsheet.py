"""
Lot 59 — lire et écrire les fichiers que les commerçants utilisent vraiment : Excel (.xlsx) et CSV.

- Lecture : la première ligne non vide est l'en-tête ; chaque cellule devient un texte propre (un numéro de
  téléphone saisi comme nombre ne devient jamais « 2.25e+12 », une date devient AAAA-MM-JJ). Formules : la
  dernière valeur calculée par Excel. Rien n'est exécuté (pas de macro).
- Écriture : un .xlsx simple qui s'ouvre directement dans Excel (en-tête en gras, figé, filtre, largeurs).
"""
import csv
import io
from datetime import date, datetime

MAX_ROWS = 10000
MAX_BYTES = 5 * 1024 * 1024


class SpreadsheetError(ValueError):
    """Fichier illisible : le message est montré tel quel."""


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "oui" if value else "non"
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value)
    return " ".join(str(value).split())


def _rows_from_csv(raw: bytes) -> list[list[str]]:
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("cp1252", errors="replace")  # Excel français enregistre souvent en Windows-1252
    sample = text[:4096]
    delimiter = ";" if sample.count(";") > sample.count(",") else ","
    return [[_cell(v) for v in row] for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def _workbook(raw: bytes):
    from openpyxl import load_workbook

    try:
        return load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 — fichier abîmé, ancien .xls, mot de passe…
        raise SpreadsheetError("Fichier Excel illisible : enregistrez-le au format .xlsx (Excel 2007 ou plus récent)") from None


def read_table(raw: bytes, filename: str, sheet: str | None = None) -> dict:
    """{sheets, sheet, header, rows, first_line} — rows sans les lignes entièrement vides."""
    name = (filename or "").lower()
    if len(raw) > MAX_BYTES:
        raise SpreadsheetError("Fichier trop volumineux (5 Mo au plus)")
    if name.endswith(".xls"):
        raise SpreadsheetError("Ancien format Excel (.xls) : enregistrez le fichier au format .xlsx")
    sheets: list[str] = []
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        book = _workbook(raw)
        sheets = list(book.sheetnames)
        chosen = sheet if sheet in sheets else sheets[0]
        table = [[_cell(v) for v in row] for row in book[chosen].iter_rows(values_only=True)]
        book.close()
        sheet = chosen
    elif name.endswith(".csv"):
        table = _rows_from_csv(raw)
        sheet = None
    else:
        raise SpreadsheetError("Format accepté : Excel (.xlsx) ou CSV")
    start = next((i for i, row in enumerate(table) if any(row)), None)
    if start is None:
        raise SpreadsheetError("Le fichier est vide")
    header = table[start]
    while header and not header[-1]:
        header = header[:-1]
    rows = []
    for row in table[start + 1:]:
        if not any(row):
            continue
        rows.append((row + [""] * len(header))[:len(header)])
        if len(rows) > MAX_ROWS:
            raise SpreadsheetError(f"Trop de lignes : {MAX_ROWS} au plus par fichier")
    return {"sheets": sheets, "sheet": sheet, "header": header, "rows": rows, "first_line": start + 2}


def to_csv_text(header: list[str], rows: list[list[str]]) -> str:
    out = io.StringIO()
    writer = csv.writer(out, delimiter=";")
    writer.writerow(header)
    writer.writerows(rows)
    return out.getvalue()


def to_xlsx(title: str, header: list[str], rows: list[list]) -> bytes:
    """Un classeur d'une feuille, prêt pour Excel."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    book = Workbook()
    ws = book.active
    ws.title = (title or "Export")[:31]
    ws.append(header)
    for row in rows:
        ws.append(list(row))
        for cell in ws[ws.max_row]:
            if cell.data_type == "f":  # un texte qui commence par « = » reste un texte, jamais une formule
                cell.data_type = "s"
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    for i, label in enumerate(header, start=1):
        width = max([len(str(label))] + [len(str(r[i - 1])) for r in rows[:500] if i - 1 < len(r) and r[i - 1] is not None])
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, width + 2))
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
