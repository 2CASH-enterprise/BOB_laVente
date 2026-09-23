import io

import pytest

from app.core.security import hash_password
from app.models.customer import Customer
from app.models.order import Order, OrderItem, OrderStatus
from app.models.product import Product
from app.models.tenant import Tenant
from app.models.user import Role, User


async def _setup(db_session, email: str, role: Role = Role.OWNER):
    tenant = Tenant(name="Boutique", country="SN", currency="XOF", email=email)
    db_session.add(tenant)
    await db_session.flush()
    db_session.add(
        User(tenant_id=tenant.id, email=email, hashed_password=hash_password("x"), full_name="U", role=role)
    )
    await db_session.commit()
    await db_session.refresh(tenant)
    return tenant


async def _login(client, email):
    r = await client.post("/api/v1/auth/login", data={"username": email, "password": "x"})
    return r.json()["access_token"]


@pytest.mark.asyncio
async def test_delete_product_soft_deletes(client, db_session, unique_email):
    tenant = await _setup(db_session, unique_email)
    product = Product(tenant_id=tenant.id, sku="X", name="X", price=1000, currency="XOF", stock_quantity=5)
    db_session.add(product)
    await db_session.commit()
    await db_session.refresh(product)
    token = await _login(client, unique_email)

    response = await client.delete(f"/api/v1/products/{product.id}", headers={"Authorization": f"Bearer {token}"})
    assert response.status_code == 204

    await db_session.refresh(product)
    assert product.active is False


@pytest.mark.asyncio
async def test_deleted_product_no_longer_in_list(client, db_session, unique_email):
    tenant = await _setup(db_session, unique_email)
    product = Product(tenant_id=tenant.id, sku="X", name="X", price=1000, currency="XOF", stock_quantity=5)
    db_session.add(product)
    await db_session.commit()
    await db_session.refresh(product)
    token = await _login(client, unique_email)
    headers = {"Authorization": f"Bearer {token}"}

    await client.delete(f"/api/v1/products/{product.id}", headers=headers)
    listing = await client.get("/api/v1/products", headers=headers)
    assert listing.json() == []


@pytest.mark.asyncio
async def test_delete_product_with_existing_orders_never_breaks_history(client, db_session, unique_email):
    """Le point critique : un produit déjà vendu doit rester supprimable sans casser l'historique."""
    tenant = await _setup(db_session, unique_email)
    product = Product(tenant_id=tenant.id, sku="X", name="X", price=1000, currency="XOF", stock_quantity=5)
    db_session.add(product)
    customer = Customer(tenant_id=tenant.id, whatsapp_number="221700000000")
    db_session.add(customer)
    await db_session.flush()
    order = Order(tenant_id=tenant.id, customer_id=customer.id, status=OrderStatus.PENDING, total_amount=1000, currency="XOF", created_by="IA")
    db_session.add(order)
    await db_session.flush()
    db_session.add(OrderItem(order_id=order.id, product_id=product.id, quantity=1, unit_price=1000, subtotal=1000))
    await db_session.commit()
    token = await _login(client, unique_email)

    response = await client.delete(f"/api/v1/products/{product.id}", headers={"Authorization": f"Bearer {token}"})
    assert response.status_code == 204  # jamais une erreur 500 de contrainte FK


@pytest.mark.asyncio
async def test_delete_requires_manager_role(client, db_session, unique_email):
    tenant = await _setup(db_session, unique_email, role=Role.AGENT)
    product = Product(tenant_id=tenant.id, sku="X", name="X", price=1000, currency="XOF", stock_quantity=5)
    db_session.add(product)
    await db_session.commit()
    await db_session.refresh(product)
    token = await _login(client, unique_email)

    response = await client.delete(f"/api/v1/products/{product.id}", headers={"Authorization": f"Bearer {token}"})
    assert response.status_code == 403


@pytest.mark.asyncio
async def test_upload_product_image_sets_absolute_url(client, db_session, unique_email):
    tenant = await _setup(db_session, unique_email)
    product = Product(tenant_id=tenant.id, sku="X", name="X", price=1000, currency="XOF", stock_quantity=5)
    db_session.add(product)
    await db_session.commit()
    await db_session.refresh(product)
    token = await _login(client, unique_email)

    fake_image = b"\xff\xd8\xff\xe0fake_jpeg_content"
    files = {"file": ("photo.jpg", io.BytesIO(fake_image), "image/jpeg")}
    response = await client.post(
        f"/api/v1/products/{product.id}/image", files=files, headers={"Authorization": f"Bearer {token}"}
    )
    assert response.status_code == 200
    assert response.json()["image_url"].startswith("https://")
    assert response.json()["image_url"].endswith(".jpg")


@pytest.mark.asyncio
async def test_upload_product_image_rejects_bad_extension(client, db_session, unique_email):
    tenant = await _setup(db_session, unique_email)
    product = Product(tenant_id=tenant.id, sku="X", name="X", price=1000, currency="XOF", stock_quantity=5)
    db_session.add(product)
    await db_session.commit()
    await db_session.refresh(product)
    token = await _login(client, unique_email)

    files = {"file": ("virus.exe", io.BytesIO(b"malicious"), "application/octet-stream")}
    response = await client.post(
        f"/api/v1/products/{product.id}/image", files=files, headers={"Authorization": f"Bearer {token}"}
    )
    assert response.status_code == 400


@pytest.mark.asyncio
async def test_export_xlsx_returns_valid_spreadsheet(client, db_session, unique_email):
    tenant = await _setup(db_session, unique_email)
    db_session.add(Product(tenant_id=tenant.id, sku="X", name="Samsung A56", price=280000, currency="XOF", stock_quantity=10))
    await db_session.commit()
    token = await _login(client, unique_email)

    response = await client.get("/api/v1/products/export/xlsx", headers={"Authorization": f"Bearer {token}"})
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(response.content))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][0] == "SKU"
    assert any(row[1] == "Samsung A56" for row in rows[1:])


@pytest.mark.asyncio
async def test_export_xlsx_isolated_by_tenant(client, db_session, unique_email):
    tenant_a = await _setup(db_session, unique_email)
    tenant_b = await _setup(db_session, f"b_{unique_email}")
    db_session.add(Product(tenant_id=tenant_a.id, sku="A", name="Produit A", price=1000, currency="XOF", stock_quantity=1))
    db_session.add(Product(tenant_id=tenant_b.id, sku="B", name="Produit B", price=1000, currency="XOF", stock_quantity=1))
    await db_session.commit()

    token_a = await _login(client, unique_email)
    response = await client.get("/api/v1/products/export/xlsx", headers={"Authorization": f"Bearer {token_a}"})

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(response.content))
    rows = list(wb.active.iter_rows(values_only=True))
    names = [row[1] for row in rows[1:]]
    assert "Produit A" in names
    assert "Produit B" not in names
