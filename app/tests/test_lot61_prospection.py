"""
Lot 61 — prospection d'AgenC'AI (Super Admin) : prospects, lien personnel suivi, entonnoir
ajouté → contacté → a cliqué → démo → compte créé → activé → payant. Aucun envoi automatique.
"""
import io
import json
import re
import uuid
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import pytest
from openpyxl import Workbook
from sqlalchemy import select

from app.core.security import hash_password
from app.models.conversation import Conversation, Message, MessageSender
from app.models.customer import Customer
from app.models.prospect import Prospect, ProspectEvent
from app.models.superadmin_user import SuperAdminUser
from app.models.tenant import Tenant
from app.models.whatsapp_account import WhatsAppAccount
from app.services import prospection as pros

ROOT = Path(__file__).resolve().parents[1]
SA_HTML = (ROOT / "static" / "superadmin" / "index.html").read_text(encoding="utf-8")
BROWSER = "Mozilla/5.0 (iPhone; CPU iPhone OS 17_0 like Mac OS X) AppleWebKit/605.1.15 Safari/604.1"
API = "/api/v1/superadmin/prospects"


async def _admin(client, db, name="Koffi Yao"):
    email = f"sa{uuid.uuid4().hex[:6]}@agencai.ci"
    db.add(SuperAdminUser(email=email, hashed_password=hash_password("supersecret123"), full_name=name))
    await db.commit()
    token = (await client.post("/api/v1/superadmin/login", data={"username": email, "password": "supersecret123"})).json()["access_token"]
    return {"Authorization": f"Bearer {token}"}


async def _prospect(client, headers, **fields):
    body = {"company": "Boutique Awa", "contact_name": "Awa Koné", "phone": "07 07 00 00 01", "sector": "ONLINE_STORE",
            "source": "Salon", **fields}
    r = await client.post(API, headers=headers, json=body)
    assert r.status_code == 201, r.text
    return r.json()


# --- Règles -------------------------------------------------------------------------------------------------

def test_bots_and_codes():
    assert pros.is_bot("WhatsApp/2.23.20.0 A") and pros.is_bot("facebookexternalhit/1.1") and pros.is_bot(None)
    assert pros.is_bot("Mozilla/5.0 (compatible; Googlebot/2.1)") and pros.is_bot("TelegramBot (like TwitterBot)")
    assert not pros.is_bot(BROWSER)
    assert pros.guess_sector("Concession automobile", "ONLINE_STORE") == "CAR_DEALERSHIP"
    assert pros.guess_sector("Courtier en assurances", "ONLINE_STORE") == "INSURANCE_BROKER"
    assert pros.guess_sector("Prêt-à-porter", "CAR_DEALERSHIP") == "ONLINE_STORE"
    assert pros.guess_sector("", "CAR_DEALERSHIP") == "CAR_DEALERSHIP"
    assert pros.guess_country("Sénégal", "CI") == "SN" and pros.guess_country("bj", "CI") == "BJ"
    assert pros.guess_country("Atlantide", "CI") == "CI"
    assert pros.detect_mapping(["Raison sociale", "Gérant", "Tél.", "E-mail", "Commune", "Activité", "Pays", "Divers"]) == [
        "company", "contact_name", "phone", "email", "city", "sector", "country", None]


def test_clean_fields():
    fields = pros.clean_fields({"company": "  Auto   Prestige ", "phone": "07 07 00 00 01", "email": " Awa@Exemple.CI ",
                                "country": "ci", "sector": "CAR_DEALERSHIP", "next_action_on": "2026-10-15"})
    assert fields["company"] == "Auto Prestige" and fields["phone"] == "2250707000001" and fields["country"] == "CI"
    assert fields["email"] == "awa@exemple.ci" and fields["next_action_on"] == date(2026, 10, 15)
    for bad, message in (({"company": " "}, "nom de l'entreprise"), ({"company": "A", "country": "XX"}, "Pays"),
                         ({"company": "A", "sector": "BANQUE"}, "Secteur"), ({"company": "A", "email": "awa@"}, "Email invalide"),
                         ({"company": "A", "phone": "12"}, "Téléphone invalide")):
        with pytest.raises(pros.ProspectError, match=message):
            pros.clean_fields(bad)


def test_stage_and_follow_up():
    now = datetime(2026, 10, 10, 12, tzinfo=timezone.utc)
    p = Prospect(company="A", status="ACTIVE")
    assert pros.stage_of(p, {}) == "ADDED"
    p.contacted_at = p.last_contact_at = now - timedelta(days=5)
    assert pros.stage_of(p, {}) == "CONTACTED" and pros.to_follow_up(p, "CONTACTED", now.date(), now)
    p.last_contact_at = now - timedelta(days=3)
    assert not pros.to_follow_up(p, "CONTACTED", now.date(), now)
    p.next_action_on = now.date()  # date de relance prévue : elle l'emporte
    assert pros.to_follow_up(p, "CONTACTED", now.date(), now)
    p.next_action_on = now.date() + timedelta(days=1)
    assert not pros.to_follow_up(p, "CONTACTED", now.date(), now)
    p.first_click_at = now
    assert pros.stage_of(p, {}) == "CLICKED"
    p.demo_at = now
    assert pros.stage_of(p, {}) == "DEMO"
    p.signed_up_at, p.tenant_id = now, uuid.uuid4()
    assert pros.stage_of(p, {}) == "SIGNED_UP" and not pros.to_follow_up(p, "SIGNED_UP", now.date(), now)
    assert pros.stage_of(p, {p.tenant_id: (True, False)}) == "ACTIVATED"
    assert pros.stage_of(p, {p.tenant_id: (True, True)}) == "PAID"
    lost = Prospect(company="B", status="NOT_INTERESTED", next_action_on=now.date())
    assert not pros.to_follow_up(lost, "CONTACTED", now.date(), now)


def test_funnel_counts_every_step_reached():
    ids = [uuid.uuid4() for _ in range(4)]
    prospects = [Prospect(id=i, company=str(n), source="Salon" if n < 3 else None, sector="ONLINE_STORE") for n, i in enumerate(ids)]
    stages = dict(zip(ids, ["ADDED", "CLICKED", "SIGNED_UP", "PAID"]))
    funnel = pros.funnel(prospects, stages)
    assert [s["count"] for s in funnel] == [4, 3, 3, 2, 2, 1, 1]
    assert [s["rate_pct"] for s in funnel] == [None, 75, 100, 67, 100, 50, 100]
    rows = pros.group_funnel(prospects, stages, lambda p: p.source)
    assert rows[0]["name"] == "Salon" and rows[0]["total"] == 3 and rows[0]["signed_up"] == 1 and rows[0]["signup_rate_pct"] == 33
    assert rows[1]["name"] == "—" and rows[1]["paid"] == 1


def test_message_and_whatsapp_link():
    p = Prospect(company="Auto <Prestige>", contact_name="Serge Kouadio", sector="CAR_DEALERSHIP", code="abc234", phone="2250707000001")
    text = pros.message(p, "Koffi")
    assert text.startswith("Bonjour Serge, je suis Koffi d'AgenC'AI.") and "rendez-vous d'essai" in text and text.endswith("/p/abc234")
    url = pros.whatsapp_url(p, "Koffi")
    assert url.startswith("https://wa.me/2250707000001?text=Bonjour%20Serge") and pros.whatsapp_url(Prospect(company="x", code="a"), None) is None
    assert pros.message(Prospect(company="X", sector="INSURANCE_BROKER", code="c"), None).startswith("Bonjour, je suis AgenC'AI.")


# --- Lien personnel --------------------------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_personal_link_counts_people_not_previews(client, db_session):
    headers = await _admin(client, db_session)
    p = await _prospect(client, headers)
    assert re.fullmatch(r"[a-z2-9]{6}", p["code"]) and p["link"].endswith(f"/p/{p['code']}")
    preview = await client.get(f"/p/{p['code']}", headers={"User-Agent": "WhatsApp/2.24"}, follow_redirects=False)
    from app.core.config import get_settings

    home = get_settings().public_base_url.rstrip("/") + "/"  # lot 62 : la présentation sous l'adresse publique de Bob
    assert preview.status_code == 302 and preview.headers["location"] == home
    stored = (await db_session.execute(select(Prospect).where(Prospect.code == p["code"]))).scalar_one()
    await db_session.refresh(stored)
    assert stored.click_count == 0 and stored.first_click_at is None  # l'aperçu de WhatsApp n'est pas un clic
    for _ in range(2):
        r = await client.get(f"/p/{p['code'].upper()}", headers={"User-Agent": BROWSER}, follow_redirects=False)
        assert r.status_code == 302 and f"{pros.COOKIE}={p['code']}" in r.headers["set-cookie"]
        assert "HttpOnly" in r.headers["set-cookie"] and "Max-Age=2592000" in r.headers["set-cookie"]
    await db_session.refresh(stored)
    assert stored.click_count == 2 and stored.first_click_at is not None
    events = (await db_session.execute(select(ProspectEvent.kind).where(ProspectEvent.prospect_id == stored.id))).scalars().all()
    assert sorted(events) == ["ADDED", "CLICK"]  # un seul événement « a cliqué »
    unknown = await client.get("/p/zzzzzz", headers={"User-Agent": BROWSER}, follow_redirects=False)
    assert unknown.status_code == 302 and "set-cookie" not in unknown.headers
    weird = await client.get("/p/%3Cscript%3E", headers={"User-Agent": BROWSER}, follow_redirects=False)
    assert weird.status_code == 302 and "set-cookie" not in weird.headers


@pytest.mark.asyncio
async def test_demo_and_signup_are_attached(client, db_session, unique_email, email_code):
    headers = await _admin(client, db_session)
    p = await _prospect(client, headers, sector="INSURANCE_BROKER", email="cabinet@exemple.ci")
    client.cookies.set(pros.COOKIE, p["code"])
    demo = await client.post("/api/v1/demo/create", data={"company_name": "Cabinet K", "business_type": "INSURANCE_BROKER", "country": "CI"})
    assert demo.status_code == 200, demo.text
    stored = (await db_session.execute(select(Prospect).where(Prospect.code == p["code"]))).scalar_one()
    await db_session.refresh(stored)
    assert stored.demo_at is not None and str(stored.demo_tenant_id) == demo.json()["tenant_id"]
    # Le compte créé depuis la démo, même sans le code (autre navigateur) : rattaché par la démo.
    client.cookies.clear()
    demo_headers = {"Authorization": f"Bearer {demo.json()['demo_token']}"}
    promoted = await client.post("/api/v1/demo/promote", headers=demo_headers, json={
        "email": unique_email, "password": "supersecret123", "full_name": "Kouassi", "verification_code": await email_code(unique_email)})
    assert promoted.status_code == 200, promoted.text
    await db_session.refresh(stored)
    assert stored.signed_up_at is not None and stored.tenant_id == stored.demo_tenant_id
    detail = (await client.get(f"{API}/{p['id']}", headers=headers)).json()
    assert detail["stage"] == "SIGNED_UP" and [e["kind"] for e in detail["events"]][:3] == ["SIGNUP", "DEMO", "ADDED"]
    # Une deuxième démo, un deuxième compte avec le même lien : la première fois reste.
    first_demo, first_account = stored.demo_tenant_id, stored.tenant_id
    client.cookies.set(pros.COOKIE, p["code"])
    again = await client.post("/api/v1/demo/create", data={"company_name": "Cabinet K2", "business_type": "INSURANCE_BROKER", "country": "CI"})
    await pros.attach(db_session, type("R", (), {"cookies": {pros.COOKIE: p["code"]}})(), uuid.UUID(again.json()["tenant_id"]), "SIGNUP")
    await db_session.refresh(stored)
    assert stored.demo_tenant_id == first_demo and stored.tenant_id == first_account
    client.cookies.clear()


@pytest.mark.asyncio
async def test_signup_attached_by_link_or_by_email(client, db_session, email_code):
    headers = await _admin(client, db_session)
    by_link = await _prospect(client, headers, company="Par le lien", phone="0707000002")
    by_email = await _prospect(client, headers, company="Par l'email", phone="0707000003", email="Gerant@Boutique.ci")
    other = await _prospect(client, headers, company="Autre", phone="0707000004")

    async def register(email):
        return await client.post("/api/v1/auth/register-tenant", json={
            "company_name": "X", "country": "CI", "currency": "XOF", "owner_email": email,
            "verification_code": await email_code(email), "owner_full_name": "A B", "owner_password": "supersecret123"})

    client.cookies.set(pros.COOKIE, by_link["code"])
    assert (await register(f"x{uuid.uuid4().hex[:6]}@l61.ci")).status_code == 201
    client.cookies.clear()
    assert (await register("gerant@boutique.ci")).status_code == 201
    assert (await register(f"y{uuid.uuid4().hex[:6]}@l61.ci")).status_code == 201  # personne : rien n'est rattaché
    stages = {r["company"]: r["stage"] for r in (await client.get(API, headers=headers)).json()["prospects"]}
    assert stages == {"Par le lien": "SIGNED_UP", "Par l'email": "SIGNED_UP", other["company"]: "ADDED"}


@pytest.mark.asyncio
async def test_attach_never_breaks_a_signup(db_session):
    class Broken:
        @property
        def cookies(self):
            raise RuntimeError("cookies illisibles")

    await pros.attach(db_session, Broken(), uuid.uuid4(), "SIGNUP")  # aucune exception


@pytest.mark.asyncio
async def test_activated_and_paid_come_from_the_shop(client, db_session):
    headers = await _admin(client, db_session)
    p = await _prospect(client, headers)
    tenant = Tenant(name="Boutique Awa", country="CI", currency="XOF", email=f"t{uuid.uuid4().hex[:6]}@l61.ci")
    db_session.add(tenant)
    await db_session.flush()
    stored = (await db_session.execute(select(Prospect).where(Prospect.code == p["code"]))).scalar_one()
    stored.signed_up_at, stored.tenant_id = datetime.now(timezone.utc), tenant.id
    db_session.add(WhatsAppAccount(tenant_id=tenant.id, waba_id="w", phone_number_id=f"pn{uuid.uuid4().hex[:6]}", system_user_token="t"))
    await db_session.commit()
    assert (await client.get(f"{API}/{p['id']}", headers=headers)).json()["stage"] == "SIGNED_UP"  # connecté, sans client
    customer = Customer(tenant_id=tenant.id, whatsapp_number="2250707000009")
    db_session.add(customer)
    await db_session.flush()
    conv = Conversation(tenant_id=tenant.id, customer_id=customer.id)
    db_session.add(conv)
    await db_session.flush()
    db_session.add(Message(tenant_id=tenant.id, conversation_id=conv.id, sender=MessageSender.CUSTOMER, content="Bonjour", message_type="text"))
    await db_session.commit()
    assert (await client.get(f"{API}/{p['id']}", headers=headers)).json()["stage"] == "ACTIVATED"
    tenant.paid_until = date(2026, 11, 10)
    await db_session.commit()
    assert (await client.get(f"{API}/{p['id']}", headers=headers)).json()["stage"] == "PAID"
    funnel = (await client.get(f"{API}/funnel", headers=headers)).json()
    assert [s["count"] for s in funnel["funnel"]] == [1] * 7  # payant : toutes les étapes comptent comme atteintes


# --- Super Admin : API --------------------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_only_superadmins(client, db_session):
    assert (await client.get(API)).status_code == 401
    shop = await client.get(API, headers={"Authorization": "Bearer pas-un-jeton"})
    assert shop.status_code == 401
    from app.tests.test_lot60_rapport import _tenant

    tenant = await _tenant(db_session)
    token = (await client.post("/api/v1/auth/login", data={"username": tenant.email, "password": "x"})).json()["access_token"]
    assert (await client.get(API, headers={"Authorization": f"Bearer {token}"})).status_code == 401  # jeton de boutique


@pytest.mark.asyncio
async def test_create_update_contact_status_note_delete(client, db_session):
    headers = await _admin(client, db_session, name="Koffi Yao")
    p = await _prospect(client, headers, email="awa@exemple.ci")
    assert p["owner"] == "Koffi Yao" and p["stage"] == "ADDED" and p["phone"] == "+2250707000001"
    assert p["whatsapp_url"].startswith("https://wa.me/2250707000001?text=Bonjour%20Awa%2C%20je%20suis%20Koffi%20d")
    dup = await client.post(API, headers=headers, json={"company": "Autre", "phone": "+225 07 07 00 00 01"})
    assert dup.status_code == 409 and "Boutique Awa" in dup.json()["detail"]
    dup = await client.post(API, headers=headers, json={"company": "Autre", "email": "AWA@exemple.ci"})
    assert dup.status_code == 409
    assert (await client.post(API, headers=headers, json={"company": "X", "email": "pas-un-email"})).status_code == 422
    # Modifier : seul ce qui est envoyé change
    other_admin = (await db_session.execute(select(SuperAdminUser).where(SuperAdminUser.full_name == "Koffi Yao"))).scalar_one()
    updated = (await client.put(f"{API}/{p['id']}", headers=headers, json={"city": "Cocody", "owner_id": str(other_admin.id)})).json()
    assert updated["city"] == "Cocody" and updated["email"] == "awa@exemple.ci" and updated["company"] == "Boutique Awa"
    assert (await client.put(f"{API}/{p['id']}", headers=headers, json={"owner_id": str(uuid.uuid4())})).status_code == 422
    assert (await client.put(f"{API}/{p['id']}", headers=headers, json={"owner_id": "n'importe quoi"})).status_code == 422
    # Contact noté à la main, puis relance prévue
    contacted = (await client.post(f"{API}/{p['id']}/contact", headers=headers,
                                   json={"channel": "PHONE", "note": "Intéressée, rappeler lundi", "next_action_on": "2026-10-12"})).json()
    assert contacted["stage"] == "CONTACTED" and contacted["next_action_on"] == "2026-10-12" and contacted["last_contact_at"]
    assert contacted["events"][0]["label"] == "Contacté" and contacted["events"][0]["channel"] == "Appel"
    assert contacted["events"][0]["actor"] == "Koffi Yao" and contacted["events"][0]["detail"] == "Intéressée, rappeler lundi"
    assert (await client.post(f"{API}/{p['id']}/contact", headers=headers, json={"channel": "PIGEON"})).status_code == 422
    noted = (await client.post(f"{API}/{p['id']}/note", headers=headers, json={"note": "A deux boutiques"})).json()
    assert noted["events"][0]["kind"] == "NOTE"
    # Ne plus contacter : plus de lien WhatsApp ni de contact possible
    stopped = (await client.post(f"{API}/{p['id']}/status", headers=headers, json={"status": "UNSUBSCRIBED", "reason": "Demande"})).json()
    assert stopped["status_label"] == "Ne plus contacter" and stopped["whatsapp_url"] is None and stopped["next_action_on"] is None
    assert (await client.post(f"{API}/{p['id']}/contact", headers=headers, json={"channel": "PHONE"})).status_code == 409
    assert (await client.post(f"{API}/{p['id']}/status", headers=headers, json={"status": "PERDU"})).status_code == 422
    # Effacer à la demande : le prospect et son historique
    assert (await client.delete(f"{API}/{p['id']}", headers=headers)).status_code == 204
    assert (await client.get(f"{API}/{p['id']}", headers=headers)).status_code == 404
    assert (await db_session.execute(select(ProspectEvent))).scalars().all() == []


@pytest.mark.asyncio
async def test_views_search_and_funnel(client, db_session):
    headers = await _admin(client, db_session)
    a = await _prospect(client, headers, company="Auto Prestige", sector="CAR_DEALERSHIP", phone="0707000011", source="Salon auto")
    b = await _prospect(client, headers, company="Boutique Fatou", phone="0707000012", city="Plateau", source="Réseau")
    c = await _prospect(client, headers, company="Cabinet Kouassi", sector="INSURANCE_BROKER", phone="0707000013", source="Réseau")
    await client.post(f"{API}/{a['id']}/contact", headers=headers, json={"channel": "VISIT", "next_action_on": "2020-01-01"})
    await client.post(f"{API}/{c['id']}/status", headers=headers, json={"status": "NOT_INTERESTED"})
    await client.get(f"/p/{b['code']}", headers={"User-Agent": BROWSER}, follow_redirects=False)
    data = (await client.get(API, headers=headers)).json()
    assert data["counts"] == {"all": 3, "todo": 1, "replied": 0, "clicked": 1, "signed_up": 0, "lost": 1}
    assert [p["company"] for p in (await client.get(f"{API}?view=todo", headers=headers)).json()["prospects"]] == ["Auto Prestige"]
    assert [p["company"] for p in (await client.get(f"{API}?view=clicked", headers=headers)).json()["prospects"]] == ["Boutique Fatou"]
    assert [p["company"] for p in (await client.get(f"{API}?view=lost", headers=headers)).json()["prospects"]] == ["Cabinet Kouassi"]
    assert [p["company"] for p in (await client.get(f"{API}?q=plateau", headers=headers)).json()["prospects"]] == ["Boutique Fatou"]
    assert [p["company"] for p in (await client.get(f"{API}?q=0707000011", headers=headers)).json()["prospects"]] == ["Auto Prestige"]
    assert [p["company"] for p in (await client.get(f"{API}?sector=INSURANCE_BROKER", headers=headers)).json()["prospects"]] == ["Cabinet Kouassi"]
    assert (await client.get(f"{API}?view=nimporte", headers=headers)).status_code == 422
    old = (await db_session.execute(select(Prospect).where(Prospect.company == "Cabinet Kouassi"))).scalar_one()
    old.created_at = datetime.now(timezone.utc) - timedelta(days=45)
    await db_session.commit()
    assert (await client.get(f"{API}/funnel?period=30", headers=headers)).json()["total"] == 2  # ajouté il y a 45 jours
    assert (await client.get(f"{API}/funnel?period=90", headers=headers)).json()["total"] == 3
    funnel = (await client.get(f"{API}/funnel?period=all", headers=headers)).json()
    assert funnel["total"] == 3 and [s["count"] for s in funnel["funnel"]][:3] == [3, 2, 1]
    assert {r["name"]: r["total"] for r in funnel["by_source"]} == {"Réseau": 2, "Salon auto": 1}
    assert {r["name"] for r in funnel["by_sector"]} == {"Commerce", "Concession automobile", "Courtier / agent d'assurance"}
    assert funnel["by_owner"][0]["total"] == 3


def _xlsx(rows) -> bytes:
    book = Workbook()
    for row in rows:
        book.active.append(row)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


@pytest.mark.asyncio
async def test_import_and_export(client, db_session):
    headers = await _admin(client, db_session, name="Awa Admin")
    await _prospect(client, headers, company="Déjà là", phone="0707000021")
    raw = _xlsx([["Raison sociale", "Gérant", "Tél.", "Email", "Activité", "Pays", "Divers"],
                 ["Auto Prestige", "Serge", 707000022, "", "Concession auto", "", "x"],
                 ["Doublon", "", "0707000021", "", "", "", ""],
                 ["Boutique Dakar", "Fatou", "77 123 45 67", "fatou@exemple.sn", "Boutique", "Sénégal", ""],
                 ["", "Sans nom", "0707000023", "", "", "", ""],
                 ["Mauvais email", "", "", "pas-un-email", "", "", ""],
                 ["Auto Prestige bis", "", "0707000022", "", "", "", ""]])
    files = {"file": ("liste.xlsx", raw, "application/octet-stream")}
    preview = (await client.post(f"{API}/import/preview", headers=headers, files=files)).json()
    assert preview["mapping"] == ["company", "contact_name", "phone", "email", "sector", "country", None] and preview["total"] == 6
    bad = await client.post(f"{API}/import", headers=headers, files=files, data={"mapping": json.dumps([None] * 7)})
    assert bad.status_code == 422 and "Entreprise" in bad.json()["detail"]
    report = (await client.post(f"{API}/import", headers=headers, files=files, data={
        "mapping": json.dumps(preview["mapping"]), "country": "CI", "sector": "ONLINE_STORE", "source": "Salon 2026"})).json()
    assert report["created"] == 2 and report["duplicates"] == 2 and report["total"] == 6
    assert [e["line"] for e in report["errors"]] == [5, 6]
    rows = {p["company"]: p for p in (await client.get(API, headers=headers)).json()["prospects"]}
    assert rows["Auto Prestige"]["sector"] == "CAR_DEALERSHIP" and rows["Auto Prestige"]["phone"] == "+2250707000022"
    assert rows["Boutique Dakar"]["country"] == "SN" and rows["Boutique Dakar"]["phone"] == "+221771234567"
    assert rows["Boutique Dakar"]["source"] == "Salon 2026" and rows["Boutique Dakar"]["owner"] == "Awa Admin"
    export = await client.get(f"{API}/export.xlsx", headers=headers)
    assert export.status_code == 200 and "prospects.xlsx" in export.headers["content-disposition"]
    from openpyxl import load_workbook

    sheet = [list(r) for r in load_workbook(io.BytesIO(export.content)).active.iter_rows(values_only=True)]
    assert sheet[0][:4] == ["Entreprise", "Contact", "Email", "Téléphone"] and len(sheet) == 4
    assert all(re.search(r"/p/[a-z2-9]{6}$", str(r[14])) for r in sheet[1:])


@pytest.mark.asyncio
async def test_purge_stale_prospects(db_session):
    old = datetime.now(timezone.utc) - timedelta(days=400)
    for name, extra in (("Oublié", {}), ("A cliqué", {"first_click_at": old}), ("Inscrit", {"tenant_id": None, "demo_at": old}),
                        ("Récent", {"created_at": datetime.now(timezone.utc)})):
        db_session.add(Prospect(code=uuid.uuid4().hex[:8], company=name, created_at=extra.pop("created_at", old), **extra))
    await db_session.commit()
    assert await pros.purge_stale(db_session) == 1
    left = sorted((await db_session.execute(select(Prospect.company))).scalars().all())
    assert left == ["A cliqué", "Inscrit", "Récent"]


def test_purge_is_scheduled():
    from app.workers.celery_app import TASK_MODULES, celery_app

    assert "app.workers.prospects" in TASK_MODULES
    assert celery_app.conf.beat_schedule["purge-prospects"]["task"] == "app.workers.prospects.purge_prospects_task"


# --- Page Super Admin -----------------------------------------------------------------------------------------

def _js(name):
    body = SA_HTML[SA_HTML.index(f"function {name}("):]
    return body[:body.index("\n}\n") + 2]


def test_superadmin_page():
    for el in ('data-section="prospects"', 'id="prospects-section"', 'id="funnel-bars"', 'id="prospect-rows"',
               'id="prospect-views"', 'id="prospect-search"', 'id="prospect-modal"', "openProspectImport()", "exportProspects()"):
        assert el in SA_HTML, el
    render = _js("renderProspect")
    for field in ("p.company", "p.link", "p.message", "ev.detail", "ev.actor", "p.notes", "p.status_reason"):
        assert f"esc({field}" in render, field
    assert "${p.company}" not in render and "${ev.detail}" not in render
    rows = _js("loadProspects")
    assert "esc(p.company)" in rows and "esc(p.contact_name" in rows and "esc(p.source" in rows
    assert 'prospectAction("contact", { channel: "WHATSAPP"' in _js("openWhatsApp")
    assert '"noopener"' in _js("openWhatsApp")
    assert "/api/v1/superadmin/prospects/import/preview" in _js("previewProspectImport")
    for name in ("prospectView", "prospectData", "funnelData", "currentProspect"):
        assert re.search(rf"^var {name} =", SA_HTML, re.M), name
