"""
Lot 59 — vos clients depuis et vers Excel (les trois secteurs) et kit de lancement (lien suivi, QR code, affiche,
textes prêts à copier).
"""
import io
import json
import re
import uuid
from datetime import datetime, timezone
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from app.models.audit_log import AuditLog
from app.models.contact_point import ContactPoint
from app.models.customer import Customer
from app.models.insurance_contract import InsuranceContract
from app.models.user import Role
from app.models.whatsapp_account import WhatsAppAccount
from app.services import customer_import, launch_kit
from app.services.business_type import CAR_DEALERSHIP, INSURANCE_BROKER, ONLINE_STORE
from app.services.insurance_contracts import normalize_phone
from app.services.plan_limits import count_contact_points
from app.services.spreadsheet import SpreadsheetError, read_table, to_xlsx
from app.tests.test_lot53_courtier import _cabinet, _headers
from app.tests.test_lot55_contrats import _login, _user

ROOT = Path(__file__).resolve().parents[1]
HTML = (ROOT / "static" / "dashboard" / "index.html").read_text(encoding="utf-8")


def _function(name):
    body = HTML[HTML.index(f"function {name}("):]
    return body[:body.index("\n}\n") + 2]


def _xlsx(rows, sheets=None) -> bytes:
    book = Workbook()
    ws = book.active
    ws.title = "Clients"
    for row in rows:
        ws.append(row)
    for name, extra in (sheets or {}).items():
        other = book.create_sheet(name)
        for row in extra:
            other.append(row)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _sheet(content: bytes) -> list[list]:
    return [list(r) for r in load_workbook(io.BytesIO(content)).active.iter_rows(values_only=True)]


# --- Lire un fichier ----------------------------------------------------------------------------------------

def test_read_xlsx_cleans_cells():
    raw = _xlsx([[], ["Tél.", "Nom", "Inscrit le", "Montant"],
                 [2250700000001, "  Awa   Koné ", datetime(2026, 3, 4), 12.5],
                 [None, None, None, None],
                 [707000002.0, "Yao", datetime(2026, 3, 4, 10, 30), True]])
    table = read_table(raw, "Clients.XLSX")
    assert table["header"] == ["Tél.", "Nom", "Inscrit le", "Montant"] and table["first_line"] == 3
    assert table["rows"] == [["2250700000001", "Awa Koné", "2026-03-04", "12.5"],  # jamais « 2.25e+12 »
                             ["707000002", "Yao", "2026-03-04 10:30", "oui"]]  # la ligne vide est sautée
    assert table["sheets"] == ["Clients"] and table["sheet"] == "Clients"


def test_numbers_stored_as_decimals():
    from app.services.spreadsheet import _cell

    assert _cell(707000002.0) == "707000002" and _cell(2.250707e12) == "2250707000000" and _cell(12.5) == "12.5"
    assert _cell(None) == "" and _cell(False) == "non"


def test_read_sheet_choice_and_csv():
    raw = _xlsx([["Téléphone"], ["1"]], sheets={"Anciens": [["Téléphone"], ["2"], ["3"]]})
    assert read_table(raw, "f.xlsx", "Anciens")["rows"] == [["2"], ["3"]]
    assert read_table(raw, "f.xlsx", "Inconnue")["sheet"] == "Clients"
    semicolon = read_table("Nom;Téléphone;\nAwa;0700000001;\nYao\n".encode("cp1252"), "f.csv")
    assert semicolon["header"] == ["Nom", "Téléphone"] and semicolon["rows"] == [["Awa", "0700000001"], ["Yao", ""]]
    comma = read_table("﻿Nom,Tel\nAwa,07\n".encode("utf-8"), "f.csv")
    assert comma["header"] == ["Nom", "Tel"] and comma["sheet"] is None


@pytest.mark.parametrize("raw,name,message", [
    (b"x", "f.xls", "Ancien format"),
    (b"x", "f.xlsx", "illisible"),
    (b"x", "f.pdf", "Format accepté"),
    (b"\n;;\n", "f.csv", "vide"),
    (b"x" * (5 * 1024 * 1024 + 1), "f.csv", "volumineux"),
])
def test_read_refusals(raw, name, message):
    with pytest.raises(SpreadsheetError, match=message):
        read_table(raw, name)


def test_too_many_rows():
    raw = ("Tel\n" + "1\n" * 10001).encode()
    with pytest.raises(SpreadsheetError, match="Trop de lignes"):
        read_table(raw, "f.csv")
    assert len(read_table(("Tel\n" + "1\n" * 10000).encode(), "f.csv")["rows"]) == 10000


def test_xlsx_writer_never_writes_a_formula():
    content = to_xlsx("Un titre bien trop long pour une feuille Excel", ["Nom", "Montant"], [["=HYPERLINK(\"x\")", 12.5]])
    book = load_workbook(io.BytesIO(content))
    ws = book.active
    assert ws.title == "Un titre bien trop long pour un" and ws.freeze_panes == "A2" and ws["A1"].font.bold
    assert ws["A2"].data_type == "s" and ws["A2"].value == '=HYPERLINK("x")' and ws["B2"].value == 12.5
    assert ws.auto_filter.ref == "A1:B2"


# --- Correspondance des colonnes ----------------------------------------------------------------------------

def test_detect_mapping():
    assert customer_import.detect_mapping(
        ["N° Tél.", "Prénoms", "NOM", "E-mail", "Commune", "Remarques", "Catégorie", "Accepte les offres", "Tel 2"]) == [
        "phone", "first_name", "last_name", "email", "city", "note", "tags", None, None]
    assert customer_import.detect_mapping(["Nom et prénoms", "Numéro WhatsApp du client"]) == ["full_name", "phone"]
    assert customer_import.detect_mapping(["Client", "Mobile"]) == ["full_name", "phone"]


def test_preview():
    raw = _xlsx([["Client", "Téléphone", "Achats"]] + [[f"C{i}", f"070000000{i}", i] for i in range(8)])
    data = customer_import.preview(raw, "liste.xlsx")
    assert data["mapping"] == ["full_name", "phone", None] and data["total"] == 8 and len(data["rows"]) == 5
    assert data["missing_phone"] is False and {"value": "phone", "label": "Téléphone"} in data["fields"]
    assert customer_import.preview(b"Nom\nAwa\n", "f.csv")["missing_phone"] is True
    with pytest.raises(customer_import.ImportError_, match="Ancien format"):
        customer_import.preview(b"x", "f.xls")


def test_mapping_checks():
    check = customer_import._check_mapping
    assert check(["A", "B"], ["phone", "inconnu"]) == ["phone", None]
    for header, mapping, message in ((["A", "B"], ["phone"], "recommencez"), (["A", "B"], "phone", "recommencez"),
                                     (["A", "B"], [None, "email"], "une seule"), (["A", "B"], ["phone", "phone"], "une seule"),
                                     (["A", "B", "C"], ["phone", "email", "email"], "colonne : Email")):
        with pytest.raises(customer_import.ImportError_, match=message):
            check(header, mapping)


def test_ivorian_and_beninese_numbers_get_their_zero_back():
    assert normalize_phone("707000001", "CI") == "2250707000001"  # Excel a mangé le 0 de 07 07 00 00 01
    assert normalize_phone("0707000001", "CI") == "2250707000001"
    assert normalize_phone("+225 07 07 00 00 01", "CI") == "2250707000001"
    assert normalize_phone("2250707000001", "CI") == "2250707000001"
    assert normalize_phone("197000001", "BJ") == "229" + "0197000001"
    assert normalize_phone("771234567", "SN") == "221771234567"  # au Sénégal, 9 chiffres sans 0


# --- Import -------------------------------------------------------------------------------------------------

async def _shop(db, business_type=ONLINE_STORE):
    return await _cabinet(db, business_type=business_type)


@pytest.mark.asyncio
async def test_import_creates_and_only_completes(db_session):
    tenant = await _shop(db_session)
    known = Customer(tenant_id=tenant.id, whatsapp_number="2250707000001", first_name="Awa", email="awa@ancien.ci",
                     tags=["VIP"], notes="Cliente fidèle")
    bare = Customer(tenant_id=tenant.id, whatsapp_number="2250707000003")
    db_session.add_all([known, bare])
    await db_session.commit()
    raw = _xlsx([
        ["Téléphone", "Prénom", "Nom", "Email", "Ville", "Note", "Catégorie", "Accepte les offres", "Achats"],
        [707000001, "Aminata", "Koné", "awa@nouveau.ci", "Cocody", "Autre note", "VIP, Grossiste", "oui", 3],
        ["+225 07 07 00 00 02", "Yao", "", "pas-un-email", "", "", "", "oui", ""],
        ["07 07 00 00 03", "", "Kouassi", "k@bob.ci", "Yopougon", "Préfère le soir", "Détail", "", ""],
        ["", "Sans", "Numéro", "", "", "", "", "", ""],
        ["12", "Trop", "Court", "", "", "", "", "", ""],
        ["0707000002", "Doublon", "", "", "Plateau", "", "", "", ""],
        ["0707000001", "", "", "", "", "", "VIP", "", ""],
    ])
    mapping = customer_import.detect_mapping(_sheet(raw)[0])
    report = await customer_import.import_customers(db_session, tenant, raw, "Mes clients.xlsx", None, mapping,
                                                    now=datetime(2026, 10, 9, tzinfo=timezone.utc))
    assert report["created"] == 1 and report["completed"] == 3 and report["unchanged"] == 1 and report["total"] == 7
    assert report["errors"] == [{"line": 5, "message": "téléphone manquant"},
                                {"line": 6, "message": "téléphone manquant ou invalide"}]
    assert report["warnings"] == [{"line": 3, "message": "email ignoré (invalide) : pas-un-email"}]
    assert report["ignored_columns"] == ["Accepte les offres", "Achats"]
    await db_session.refresh(known)
    # Rien de connu n'est remplacé : prénom, email et note restent ; la ville manquait, l'étiquette s'ajoute.
    assert (known.first_name, known.last_name, known.email, known.notes) == ("Awa", None, "awa@ancien.ci", "Cliente fidèle")
    assert known.city == "Cocody" and known.tags == ["VIP", "Grossiste"]
    assert known.marketing_consent is not True  # l'accord pour les offres ne s'importe jamais
    await db_session.refresh(bare)
    assert (bare.first_name, bare.last_name, bare.email, bare.city, bare.notes, bare.tags) == (
        None, "Kouassi", "k@bob.ci", "Yopougon", "Préfère le soir", ["Détail"])
    assert bare.email_source == "MANUAL" and bare.acquisition_source is None  # sa source d'origine reste
    yao = (await db_session.execute(select(Customer).where(Customer.whatsapp_number == "2250707000002"))).scalar_one()
    assert (yao.first_name, yao.email, yao.city) == ("Yao", None, "Plateau")  # la ligne en double a complété la ville
    assert yao.acquisition_source == "IMPORT" and yao.acquisition_detail == "Fichier Mes clients.xlsx"
    assert not yao.marketing_consent
    again = await customer_import.import_customers(db_session, tenant, raw, "Mes clients.xlsx", None, mapping)
    assert again["created"] == 0 and again["completed"] == 0 and again["unchanged"] == 5
    assert (await db_session.execute(select(func.count(Customer.id)).where(Customer.tenant_id == tenant.id))).scalar_one() == 3


@pytest.mark.asyncio
async def test_full_name_and_isolation(db_session):
    tenant = await _shop(db_session, CAR_DEALERSHIP)
    other = await _shop(db_session)
    db_session.add(Customer(tenant_id=other.id, whatsapp_number="2250707000009", first_name="Autre"))
    await db_session.commit()
    raw = "Nom et prénoms;Numéro\nKonan Brou Serge;0707000009\n".encode()
    report = await customer_import.import_customers(db_session, tenant, raw, "f.csv", None, ["full_name", "phone"])
    assert report["created"] == 1  # le même numéro chez une autre boutique ne compte pas
    mine = (await db_session.execute(select(Customer).where(Customer.tenant_id == tenant.id))).scalar_one()
    assert (mine.first_name, mine.last_name) == ("Konan Brou Serge", None)
    theirs = (await db_session.execute(select(Customer).where(Customer.tenant_id == other.id))).scalar_one()
    assert theirs.first_name == "Autre"


@pytest.mark.asyncio
async def test_import_api_rights_and_audit(client, db_session):
    tenant = await _shop(db_session)
    owner = await _headers(client, tenant)
    agent = await _login(client, await _user(db_session, tenant, Role.AGENT))
    manager = await _login(client, await _user(db_session, tenant, Role.MANAGER))
    raw = _xlsx([["Tel", "Prénom"], ["0707000001", "Awa"]])
    files = {"file": ("clients.xlsx", raw, "application/octet-stream")}
    assert (await client.post("/api/v1/customers/import/preview", headers=agent, files=files)).status_code == 403
    assert (await client.post("/api/v1/customers/import", headers=agent, files=files, data={"mapping": "[]"})).status_code == 403
    preview = (await client.post("/api/v1/customers/import/preview", headers=manager, files=files)).json()
    assert preview["mapping"] == ["phone", "first_name"] and preview["total"] == 1
    bad = await client.post("/api/v1/customers/import/preview", headers=manager, files={"file": ("f.xls", b"x", "x")})
    assert bad.status_code == 422 and "xlsx" in bad.json()["detail"]
    unreadable = await client.post("/api/v1/customers/import", headers=manager, files=files, data={"mapping": "{pas du json"})
    assert unreadable.status_code == 422
    no_phone = await client.post("/api/v1/customers/import", headers=manager, files=files, data={"mapping": json.dumps([None, "first_name"])})
    assert no_phone.status_code == 422 and "Téléphone" in no_phone.json()["detail"]
    done = await client.post("/api/v1/customers/import", headers=owner, files=files,
                             data={"mapping": json.dumps(preview["mapping"]), "sheet": "Clients"})
    assert done.status_code == 200 and done.json()["created"] == 1
    actions = set((await db_session.execute(select(AuditLog.action))).scalars().all())
    assert "CUSTOMERS_IMPORTED" in actions
    big = await client.post("/api/v1/customers/import/preview", headers=manager,
                            files={"file": ("f.csv", b"x" * (5 * 1024 * 1024 + 1), "text/csv")})
    assert big.status_code == 413


# --- Export -------------------------------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_customer_export(client, db_session):
    tenant = await _shop(db_session)
    db_session.add_all([
        Customer(tenant_id=tenant.id, whatsapp_number="2250707000001", first_name="Awa", last_name="Koné", city="Cocody",
                 marketing_consent=True, marketing_consent_given_at=datetime(2026, 10, 1, 9, 0, tzinfo=timezone.utc),
                 acquisition_source="IMPORT", acquisition_detail="Fichier x.xlsx", tags=["VIP", "Gros"], notes="=1+1"),
        Customer(tenant_id=tenant.id, whatsapp_number="2250707000002"),
    ])
    other = await _shop(db_session)
    db_session.add(Customer(tenant_id=other.id, whatsapp_number="2250707000003", first_name="Ailleurs"))
    await db_session.commit()
    agent = await _login(client, await _user(db_session, tenant, Role.AGENT))
    viewer = await _login(client, await _user(db_session, tenant, Role.VIEWER))
    assert (await client.get("/api/v1/customers/export.xlsx", headers=viewer)).status_code == 403
    response = await client.get("/api/v1/customers/export.xlsx", headers=agent)
    assert response.status_code == 200 and "clients.xlsx" in response.headers["content-disposition"]
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    rows = _sheet(response.content)
    assert rows[0] == ["Prénom", "Nom", "Téléphone", "Email", "Ville", "Source", "Détail de la source", "Accepte les offres",
                       "Accord donné le", "Dernier échange", "Client depuis", "Étiquettes", "Note"]  # commerce : pas de score
    assert len(rows) == 3 and "Ailleurs" not in str(rows)
    awa = rows[1]
    assert awa[:5] == ["Awa", "Koné", "+2250707000001", None, "Cocody"] and awa[5] == "Import"
    assert awa[7:9] == ["oui", "01/10/2026 09:00"] and awa[11:] == ["VIP, Gros", "=1+1"]  # heure d'Abidjan ; jamais une formule
    assert rows[2][7:9] == ["non", None]
    tenant.country = "FR"  # les heures suivent le pays de la boutique (Paris en octobre : UTC+2)
    await db_session.commit()
    _, paris = await customer_import.export_rows(db_session, tenant)
    assert paris[0][8] == "01/10/2026 11:00"
    assert "CUSTOMERS_EXPORTED" in set((await db_session.execute(select(AuditLog.action))).scalars().all())


@pytest.mark.asyncio
async def test_export_score_column_for_broker_and_dealer(db_session):
    for sector in (INSURANCE_BROKER, CAR_DEALERSHIP):
        tenant = await _shop(db_session, sector)
        db_session.add(Customer(tenant_id=tenant.id, whatsapp_number=f"2250{uuid.uuid4().int % 10**9:09d}"))
        await db_session.commit()
        header, rows = await customer_import.export_rows(db_session, tenant)
        assert header[10] == "Score" and header[11] == "Client depuis" and len(rows[0]) == len(header)


# --- Registre des contrats : import Excel ---------------------------------------------------------------------

@pytest.mark.asyncio
async def test_contract_import_accepts_xlsx(client, db_session):
    tenant = await _cabinet(db_session)
    headers = await _headers(client, tenant)
    raw = _xlsx([["Téléphone", "Branche", "Échéance", "Assureur"], [707000001, "Auto", datetime(2027, 1, 15), "Sunu"]])
    report = (await client.post("/api/v1/contracts/import-csv", headers=headers,
                                files={"file": ("registre.xlsx", raw, "application/octet-stream")})).json()
    assert report["created"] == 1 and report["customers_created"] == 1
    contract = (await db_session.execute(select(InsuranceContract))).scalar_one()
    assert str(contract.expires_on) == "2027-01-15"
    customer = (await db_session.execute(select(Customer).where(Customer.tenant_id == tenant.id))).scalar_one()
    assert customer.whatsapp_number == "2250707000001"
    broken = await client.post("/api/v1/contracts/import-csv", headers=headers, files={"file": ("r.xlsx", b"x", "x")})
    assert broken.status_code == 422 and "illisible" in broken.json()["detail"]


# --- Kit de lancement ---------------------------------------------------------------------------------------

async def _connected(db, business_type=ONLINE_STORE, number="+225 07 07 00 00 99"):
    tenant = await _cabinet(db, business_type=business_type)
    account = (await db.execute(select(WhatsAppAccount).where(WhatsAppAccount.tenant_id == tenant.id))).scalar_one()
    account.display_phone_number = number
    await db.commit()
    return tenant


@pytest.mark.asyncio
async def test_no_kit_without_whatsapp(client, db_session):
    tenant = await _cabinet(db_session, with_account=False)
    headers = await _headers(client, tenant)
    assert (await client.get("/api/v1/kit", headers=headers)).json() == {"connected": False}
    assert (await client.put("/api/v1/kit", headers=headers, json={"greeting": "Bonjour"})).status_code == 409
    assert (await client.get("/api/v1/kit/poster", headers=headers)).status_code == 409
    assert (await db_session.execute(select(func.count(ContactPoint.id)))).scalar_one() == 0
    unnumbered = await _cabinet(db_session)  # compte WhatsApp sans numéro affiché
    assert (await client.get("/api/v1/kit", headers=await _headers(client, unnumbered))).json() == {"connected": False}


@pytest.mark.asyncio
async def test_kit_link_created_once_and_outside_plan_limit(client, db_session):
    tenant = await _connected(db_session)
    headers = await _headers(client, tenant)
    first = (await client.get("/api/v1/kit", headers=headers)).json()
    second = (await client.get("/api/v1/kit", headers=headers)).json()
    assert first["connected"] and first["link"] == second["link"] and re.search(r"/w/KIT_[0-9a-f]{8}$", first["link"])
    assert first["direct_link"] == "https://wa.me/2250707000099" and first["phone"] == "+225 07 07 00 00 99"
    assert first["greeting"] == launch_kit.GREETINGS[ONLINE_STORE] and first["active"] is True
    assert first["clicks"] == 0 and first["customers"] == 0
    assert first["qr_svg"].startswith("<svg") and "QR code WhatsApp" in first["qr_svg"]
    points = (await db_session.execute(select(ContactPoint).where(ContactPoint.tenant_id == tenant.id))).scalars().all()
    assert len(points) == 1 and points[0].channel == "PRINT" and points[0].name == launch_kit.KIT_NAME
    assert await count_contact_points(db_session, tenant.id) == 0  # le lien du kit ne compte pas dans le plan
    db_session.add(ContactPoint(tenant_id=tenant.id, code="VITRINE", name="Vitrine", greeting="Bonjour", channel="PRINT", active=True))
    await db_session.commit()
    assert await count_contact_points(db_session, tenant.id) == 1
    vitrine = (await db_session.execute(select(ContactPoint).where(ContactPoint.code == "VITRINE"))).scalar_one()
    db_session.add_all([
        Customer(tenant_id=tenant.id, whatsapp_number="2250707000001", acquisition_source="LINK", acquisition_contact_point_id=points[0].id),
        Customer(tenant_id=tenant.id, whatsapp_number="2250707000002", acquisition_source="LINK", acquisition_contact_point_id=vitrine.id),
    ])
    await db_session.commit()
    assert (await client.get("/api/v1/kit", headers=headers)).json()["customers"] == 1  # seulement ceux venus par le kit
    points[0].archived_at = datetime.now(timezone.utc)  # archivé depuis Intégrations : le kit en recrée un
    await db_session.commit()
    renewed = (await client.get("/api/v1/kit", headers=headers)).json()
    assert renewed["link"] != first["link"] and renewed["customers"] == 0


@pytest.mark.asyncio
async def test_kit_greeting(client, db_session):
    tenant = await _connected(db_session, INSURANCE_BROKER)
    owner = await _headers(client, tenant)
    agent = await _login(client, await _user(db_session, tenant, Role.AGENT))
    assert (await client.get("/api/v1/kit", headers=agent)).json()["greeting"] == launch_kit.GREETINGS[INSURANCE_BROKER]
    assert (await client.put("/api/v1/kit", headers=agent, json={"greeting": "Salut"})).status_code == 403
    assert (await client.put("/api/v1/kit", headers=owner, json={"greeting": "   "})).status_code == 422
    assert (await client.put("/api/v1/kit", headers=owner, json={"greeting": "x" * 301})).status_code == 422
    saved = (await client.put("/api/v1/kit", headers=owner, json={"greeting": "  Bonjour,   un devis   auto  "})).json()
    assert saved["greeting"] == "Bonjour, un devis auto"
    assert "KIT_GREETING_UPDATED" in set((await db_session.execute(select(AuditLog.action))).scalars().all())


@pytest.mark.asyncio
async def test_kit_texts_by_sector():
    from app.models.tenant import Tenant

    link = "https://bob.example/w/KIT_1"
    shop_tu = Tenant(name="Chez Awa", business_type=ONLINE_STORE, address_form="TU")
    shop_vous = Tenant(name="Chez Awa", business_type=ONLINE_STORE, address_form="VOUS")
    broker = Tenant(name="Cabinet K", business_type=INSURANCE_BROKER, address_form="TU")  # le courtier vouvoie toujours
    dealer = Tenant(name="Auto CI", business_type=CAR_DEALERSHIP)
    tu = launch_kit.texts(shop_tu, link)
    assert len(tu) == 5 and all(link in t["text"] for t in tu)
    assert all("tes " not in t["title"] and "ton " not in t["title"] for t in tu)  # les titres s'adressent au commerçant
    assert "Écris-nous" in tu[0]["text"] and "vous" not in " ".join(t["text"] for t in tu).lower().replace("nous", "")
    for tenant in (shop_vous, broker, dealer):
        texts = launch_kit.texts(tenant, link)
        joined = " ".join(t["text"] for t in texts)
        assert "Écrivez-nous" in joined and "Écris-nous" not in joined and " te " not in joined
    assert "assurance" in launch_kit.texts(broker, link)[0]["text"] and "renouvellement" in launch_kit.texts(broker, link)[2]["text"]
    assert "véhicule" in launch_kit.texts(dealer, link)[0]["text"] and "reprise" in launch_kit.texts(dealer, link)[2]["text"]
    assert "Cabinet K" in launch_kit.texts(broker, link)[3]["text"]


@pytest.mark.asyncio
async def test_poster_is_escaped_and_printable(client, db_session):
    tenant = await _connected(db_session, CAR_DEALERSHIP)
    tenant.name = "<script>alert(1)</script> Auto"
    await db_session.commit()
    response = await client.get("/api/v1/kit/poster", headers=await _headers(client, tenant))
    assert response.status_code == 200 and response.headers["cache-control"] == "no-store"
    page = response.text
    assert "<script>alert" not in page and "&lt;script&gt;alert(1)&lt;/script&gt; Auto" in page
    assert "Écrivez-nous sur WhatsApp" in page and "+225 07 07 00 00 99" in page and "<svg" in page
    assert "Disponibilité, essai" in page and "@page" in page and "window.print()" in page


# --- Tableau de bord ----------------------------------------------------------------------------------------

def test_dashboard():
    for el in ('id="kit-card"', 'id="kit-content"', 'id="customers-import-card"', 'id="customers-drop"',
               'id="customers-file-input"', 'id="customers-import-preview"', 'id="customers-import-msg"',
               "downloadExport('/api/v1/customers/export.xlsx', 'clients.xlsx')"):
        assert el in HTML, el
    assert 'accept=".xlsx,.csv' in HTML
    for name in ("loadKit", "copyText", "saveKitGreeting", "openKitPoster", "downloadKitQr", "initCustomersDrop",
                 "previewCustomerImport", "renderImportPreview", "checkImportMapping", "runCustomerImport"):
        assert f"function {name}(" in HTML, name
    assert "/api/v1/customers/import/preview" in _function("previewCustomerImport")
    run = _function("runCustomerImport")
    assert "/api/v1/customers/import" in run and "mapping" in run
    assert "Authorization" in _function("openKitPoster")
    show = HTML[HTML.index("function showTab("):]
    show = show[:show.index("\n}\n")]
    assert "loadKit()" in show and "initCustomersDrop()" in show
    assert "downloadExport(\"/api/v1/complaints/export.xlsx\"" in HTML
    assert "export.csv" not in HTML  # tous les exports sont en Excel
    assert re.search(r"/\\\.\(xlsx\|csv\)\$/i", HTML)
