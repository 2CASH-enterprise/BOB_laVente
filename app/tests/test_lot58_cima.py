"""
Lot 58 — courtier, règlement CIMA 01-24 : exports (art. 4, 5, 14, 15), identité de la compagnie (art. 12),
information sur les données personnelles (art. 7) et lien vers les conditions tarifaires (art. 11).
"""
import uuid
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path

import pytest
from sqlalchemy import select

from app.agents.prompts import build_system_prompt, insurance_office_info
from app.models.audit_log import AuditLog
from app.models.conversation import Conversation, ConversationStatus, Message, MessageSender
from app.models.customer import Customer
from app.models.insurance_complaint import InsuranceComplaint
from app.models.insurance_contract import InsuranceContract
from app.models.quote_request import QuoteRequest
from app.models.user import Role
from app.services import insurance
from app.services.business_type import CAR_DEALERSHIP
from app.tests.fakes import text_response
from app.tests.test_handoff_rules import wire  # noqa: F401 — fixture
from app.tests.test_lot53_courtier import _cabinet, _headers, _payload
from app.tests.test_lot55_contrats import _login, _user

ROOT = Path(__file__).resolve().parents[1]
HTML = (ROOT / "static" / "dashboard" / "index.html").read_text(encoding="utf-8")
PROFILE = "/api/v1/tenants/me/insurance-profile"


def _function(name):
    body = HTML[HTML.index(f"function {name}("):]
    return body[:body.index("\n}\n") + 2]


# --- Liens et compagnies : règles -------------------------------------------------------------------------

def test_clean_url_and_partners():
    assert insurance.clean_url("https://cabinet.ci/tarifs") == "https://cabinet.ci/tarifs"
    assert insurance.clean_url("  cabinet.ci/tarifs-2026.pdf ") == "https://cabinet.ci/tarifs-2026.pdf"
    assert insurance.clean_url("") is None and insurance.clean_url(None) is None
    for bad in ("pas un lien", "https://", "javascript:alert(1)", "https://x" + "a" * 300 + ".ci", "https://a b.ci"):
        with pytest.raises(ValueError):
            insurance.clean_url(bad)
    assert insurance.parse_partners([{"name": "  Sunu   Assurances ", "address": " Plateau "}, {"name": ""}, "x"]) == [
        {"name": "Sunu Assurances", "address": "Plateau"}]
    assert insurance.parse_partners([{"name": "NSIA"}]) == [{"name": "NSIA", "address": None}]
    with pytest.raises(ValueError):
        insurance.parse_partners([{"name": f"C{i}"} for i in range(21)])


def test_links_are_never_amounts():
    assert not insurance.contains_amount("Nos conditions : https://cabinet.ci/tarifs/2026/1.500.000-auto.pdf")
    assert insurance.contains_amount("La prime est de 85 000 FCFA, voir https://cabinet.ci")


@pytest.mark.asyncio
async def test_profile_api(client, db_session):
    tenant = await _cabinet(db_session)
    headers = await _headers(client, tenant)
    saved = await client.put(PROFILE, headers=headers, json={
        "insurance_structure": "AGENT", "insurer_name": "Sunu", "insurer_legal_name": "Sunu Assurances IARD Côte d'Ivoire",
        "insurer_address": "Plateau, Abidjan", "tariff_url": "cabinet.ci/tarifs", "privacy_policy_url": "https://cabinet.ci/vie-privee",
        "insurance_partners": [{"name": "Ignoré pour un agent"}]})
    assert saved.status_code == 200, saved.text
    body = saved.json()
    assert body["insurer_legal_name"] == "Sunu Assurances IARD Côte d'Ivoire" and body["insurer_address"] == "Plateau, Abidjan"
    assert body["tariff_url"] == "https://cabinet.ci/tarifs" and body["privacy_policy_url"] == "https://cabinet.ci/vie-privee"
    assert body["insurance_partners"] == []  # un agent a une seule compagnie
    # Un champ absent ne change rien (anciens écrans) ; vide l'efface
    kept = (await client.put(PROFILE, headers=headers, json={"insurance_structure": "AGENT", "insurer_name": "Sunu"})).json()
    assert kept["tariff_url"] == "https://cabinet.ci/tarifs" and kept["insurer_legal_name"]
    cleared = (await client.put(PROFILE, headers=headers, json={"insurance_structure": "AGENT", "tariff_url": ""})).json()
    assert cleared["tariff_url"] is None
    courtier = (await client.put(PROFILE, headers=headers, json={
        "insurance_structure": "COURTIER", "insurance_partners": [{"name": "NSIA", "address": "Cocody"}, {"name": "Allianz"}]})).json()
    assert courtier["insurance_partners"] == [{"name": "NSIA", "address": "Cocody"}, {"name": "Allianz", "address": None}]
    assert courtier["insurer_legal_name"] is None and courtier["insurer_address"] is None and courtier["insurer_name"] is None
    bad = await client.put(PROFILE, headers=headers, json={"insurance_structure": "COURTIER", "privacy_policy_url": "nimporte quoi"})
    assert bad.status_code == 422 and "Lien invalide" in bad.json()["detail"]
    manager = await _login(client, await _user(db_session, tenant, Role.MANAGER))
    assert (await client.put(PROFILE, headers=manager, json={"insurance_structure": "COURTIER"})).status_code == 403


@pytest.mark.asyncio
async def test_bob_knows_company_links_and_partners(db_session):
    tenant = await _cabinet(db_session)
    tenant.insurance_structure, tenant.insurer_name = "AGENCE_GENERALE", "Sunu"
    tenant.insurer_legal_name, tenant.insurer_address = "Sunu Assurances IARD", "Plateau, Abidjan"
    tenant.tariff_url, tenant.privacy_policy_url = "https://cabinet.ci/tarifs", "https://cabinet.ci/vie-privee"
    info = insurance_office_info(tenant)
    for line in ("Raison sociale de la compagnie : Sunu Assurances IARD", "Adresse de la compagnie : Plateau, Abidjan",
                 "Conditions tarifaires (lien) : https://cabinet.ci/tarifs",
                 "Politique de confidentialité (lien) : https://cabinet.ci/vie-privee"):
        assert line in info, line
    tenant.insurance_structure = "COURTIER"
    tenant.insurance_partners = [{"name": "NSIA", "address": "Cocody"}, {"name": "Allianz", "address": None}]
    info = insurance_office_info(tenant)
    assert "Compagnies partenaires : NSIA (Cocody) ; Allianz" in info and "Raison sociale de la compagnie" not in info
    prompt = build_system_prompt(tenant, [], "")
    assert "lien vers les conditions" in prompt and "sans en citer aucun chiffre" in prompt
    from app.services.handoff_rules import _insurance_rule

    assert "lien vers les conditions tarifaires" in _insurance_rule({"DEMANDE_REMISE"}, set()).instruction


# --- Données personnelles : mention unique -----------------------------------------------------------------

@pytest.mark.asyncio
async def test_privacy_notice_once_on_first_reply(client, db_session, wire):  # noqa: F811
    tenant = await _cabinet(db_session)
    tenant.privacy_policy_url = "https://cabinet.ci/vie-privee"
    await db_session.commit()
    wire({"intents": ["INFO_PRODUIT"], "objections": []}, [text_response("Bonjour ! Quelle assurance ?"),
                                                           text_response("Très bien.")])
    first = (await client.post("/webhooks/whatsapp", json=_payload(tenant.pnid, "2250711111111", "Bonjour"))).json()
    assert first["ai_reply"] == ("Bonjour ! Quelle assurance ?\n\nVos informations sont traitées par Cabinet Kouassi pour "
                                 "répondre à votre demande. En savoir plus : https://cabinet.ci/vie-privee")
    second = (await client.post("/webhooks/whatsapp", json=_payload(tenant.pnid, "2250711111111", "Auto"))).json()
    assert second["ai_reply"] == "Très bien."
    customer = (await db_session.execute(select(Customer).execution_options(populate_existing=True))).scalar_one()
    assert customer.privacy_notice_at is not None


@pytest.mark.asyncio
async def test_no_privacy_notice_without_link_or_outside_insurance(client, db_session, wire):  # noqa: F811
    tenant = await _cabinet(db_session)
    wire({"intents": ["INFO_PRODUIT"], "objections": []}, [text_response("Bonjour !")])
    reply = (await client.post("/webhooks/whatsapp", json=_payload(tenant.pnid, "2250722222222", "Bonjour"))).json()
    assert reply["ai_reply"] == "Bonjour !"
    dealer = await _cabinet(db_session, business_type=CAR_DEALERSHIP)
    dealer.privacy_policy_url = "https://garage.ci/vie-privee"
    await db_session.commit()
    wire({"intents": ["INFO_PRODUIT"], "objections": []}, [text_response("Bonjour !")])
    reply = (await client.post("/webhooks/whatsapp", json=_payload(dealer.pnid, "2250733333333", "Bonjour"))).json()
    assert reply["ai_reply"] == "Bonjour !"
    assert all(c.privacy_notice_at is None for c in (await db_session.execute(
        select(Customer).execution_options(populate_existing=True))).scalars().all())


# --- Exports -----------------------------------------------------------------------------------------------

async def _data(db, tenant):
    customer = Customer(tenant_id=tenant.id, whatsapp_number="2250700000011", first_name="Awa", last_name="Koné",
                        email="awa@exemple.ci", marketing_consent=True,
                        marketing_consent_given_at=datetime(2026, 10, 1, 9, tzinfo=timezone.utc))
    db.add(customer)
    await db.flush()
    conversation = Conversation(tenant_id=tenant.id, customer_id=customer.id, status=ConversationStatus.ACTIVE)
    db.add(conversation)
    await db.flush()
    db.add_all([
        Message(tenant_id=tenant.id, conversation_id=conversation.id, sender=MessageSender.CUSTOMER,
                content="<script>alert('x')</script> Bonjour"),
        Message(tenant_id=tenant.id, conversation_id=conversation.id, sender=MessageSender.AI, content="Bonjour Awa !"),
        InsuranceContract(tenant_id=tenant.id, customer_id=customer.id, branch="AUTO", insurer="SUNU", policy_number="POL-9",
                          premium=Decimal("185000"), currency="XOF", expires_on=date(2026, 12, 31), term="ANNUEL", status="ACTIVE"),
        QuoteRequest(tenant_id=tenant.id, customer_id=customer.id, branch="SANTE", client_type="PARTICULIER",
                     details={"persons_count": 3, "risk_country": "SN"}, status="LOST", lost_reason="Trop cher",
                     consent_at=datetime(2026, 10, 2, 9, tzinfo=timezone.utc), submitted_at=datetime(2026, 10, 2, 9, 1, tzinfo=timezone.utc),
                     closed_at=datetime(2026, 10, 5, 9, tzinfo=timezone.utc)),
        InsuranceComplaint(tenant_id=tenant.id, customer_id=customer.id, reference="R-2026-0001", kind="RECLAMATION",
                           channel="WHATSAPP", subject="Personne ne me rappelle", status="RESOLVED",
                           received_at=datetime(2026, 10, 3, 9, tzinfo=timezone.utc), resolution="Rappelé le jour même",
                           resolved_at=datetime(2026, 10, 3, 15, tzinfo=timezone.utc), created_by="BOB"),
    ])
    await db.commit()
    return customer


def _sheet(content: bytes) -> list[list]:
    from io import BytesIO

    from openpyxl import load_workbook

    return [list(r) for r in load_workbook(BytesIO(content)).active.iter_rows(values_only=True)]


@pytest.mark.asyncio
async def test_contract_and_quote_exports(client, db_session):
    tenant = await _cabinet(db_session)
    await _data(db_session, tenant)
    owner = await _headers(client, tenant)
    agent = await _login(client, await _user(db_session, tenant, Role.AGENT))
    full = (await client.get("/api/v1/contracts/export.xlsx", headers=owner))  # lot 59 : Excel
    assert full.status_code == 200 and "registre_contrats.xlsx" in full.headers["content-disposition"]
    lines = _sheet(full.content)
    assert lines[0][7:9] == ["Prime", "Devise"] and lines[1][7:9] == [185000, "XOF"] and "POL-9" in lines[1]
    hidden = _sheet((await client.get("/api/v1/contracts/export.xlsx", headers=agent)).content)
    assert "Prime" not in hidden[0] and 185000 not in hidden[1]
    quotes = _sheet((await client.get("/api/v1/quote-requests/export.xlsx", headers=agent)).content)
    assert quotes[0][:3] == ["Client", "Téléphone", "Branche"] and len(quotes) == 2
    row = " | ".join(str(v) for v in quotes[1] if v is not None)
    assert "Assurance santé" in row and "Perdu" in row and "Trop cher" in row and "Sénégal" in row
    assert "Nombre de personnes : 3" in row and "Risque situé hors de Côte d'Ivoire" in row
    assert "02/10/2026 09:00" in row and "02/10/2026 09:01" in row  # heure d'Abidjan
    tenant.country = "FR"  # heure de Paris en octobre : UTC+2 — les heures suivent le pays du cabinet
    await db_session.commit()
    paris = " | ".join(str(v) for v in _sheet((await client.get("/api/v1/quote-requests/export.xlsx", headers=agent)).content)[1] if v)
    assert "02/10/2026 11:00" in paris and "02/10/2026 11:01" in paris
    viewer = await _login(client, await _user(db_session, tenant, Role.VIEWER))
    assert (await client.get("/api/v1/contracts/export.xlsx", headers=viewer)).status_code == 403
    assert (await client.get("/api/v1/quote-requests/export.xlsx", headers=viewer)).status_code == 403
    actions = set((await db_session.execute(select(AuditLog.action))).scalars().all())
    assert {"CONTRACTS_EXPORTED", "QUOTE_REQUESTS_EXPORTED"} <= actions


@pytest.mark.asyncio
async def test_customer_dossier(client, db_session):
    tenant = await _cabinet(db_session)
    customer = await _data(db_session, tenant)
    owner = await _headers(client, tenant)
    page = await client.get(f"/api/v1/customers/{customer.id}/dossier", headers=owner)
    assert page.status_code == 200 and page.headers["content-type"].startswith("text/html")
    assert page.headers["cache-control"] == "no-store"
    html = page.text
    for piece in ("Dossier du client — Awa Koné (+2250700000011)", "Identité et accords", "oui, le 01/10/2026 09:00",
                  "Demandes de cotation (1)", "Perdu : 05/10/2026", "Trop cher", "Risque situé hors de Côte d&#x27;Ivoire",
                  "Contrats (1)", "POL-9", "185 000 XOF", "Réclamations et sinistres (1)", "R-2026-0001", "Rappelé le jour même",
                  "Conversations (1)", "Bonjour Awa !", "window.print()", "Imprimer / Enregistrer en PDF"):
        assert piece in html, piece
    assert "<script>alert" not in html and "&lt;script&gt;alert(&#x27;x&#x27;)&lt;/script&gt;" in html  # tout est échappé
    customer.notes = '<img src=x onerror="alert(1)">'
    await db_session.commit()
    noted = (await client.get(f"/api/v1/customers/{customer.id}/dossier", headers=owner)).text
    assert "<img src=x" not in noted and "&lt;img src=x onerror=&quot;alert(1)&quot;&gt;" in noted
    agent = await _login(client, await _user(db_session, tenant, Role.AGENT))
    assert "185 000" not in (await client.get(f"/api/v1/customers/{customer.id}/dossier", headers=agent)).text
    other = await _cabinet(db_session)
    assert (await client.get(f"/api/v1/customers/{customer.id}/dossier", headers=await _headers(client, other))).status_code == 404
    dealer = await _cabinet(db_session, business_type=CAR_DEALERSHIP)
    assert (await client.get(f"/api/v1/customers/{customer.id}/dossier",
                             headers=await _headers(client, dealer))).status_code == 403
    viewer = await _login(client, await _user(db_session, tenant, Role.VIEWER))
    assert (await client.get(f"/api/v1/customers/{customer.id}/dossier", headers=viewer)).status_code == 403
    assert "CUSTOMER_DOSSIER_EXPORTED" in set((await db_session.execute(select(AuditLog.action))).scalars().all())


# --- Tableau de bord ----------------------------------------------------------------------------------------

def test_dashboard():
    for el in ('id="insurer-legal-name"', 'id="insurer-address"', 'id="insurance-partners"', 'id="tariff-url"', 'id="privacy-url"',
               'id="insurance-partners-row"'):
        assert el in HTML, el
    save = _function("saveInsuranceProfile")
    for key in ("insurer_legal_name", "insurer_address", "insurance_partners: parsePartners(", "tariff_url", "privacy_policy_url"):
        assert key in save, key
    assert 'classList.toggle("hidden", chosen !== "COURTIER")' in _function("insuranceStructureChanged")
    assert "downloadExport('/api/v1/contracts/export.xlsx'" in HTML and "downloadExport('/api/v1/quote-requests/export.xlsx'" in HTML
    assert "openCustomerDossier('${esc(c.id)}')" in HTML and 'currentBusinessType === "INSURANCE_BROKER"' in HTML
    dossier = _function("openCustomerDossier")
    assert "Authorization" in dossier and "token" not in dossier.split("fetch(")[1].split(",")[0]  # jamais dans l'adresse
    assert "Authorization" in _function("downloadExport")


def test_partner_parsing_runs():
    import json
    import shutil
    import subprocess

    if shutil.which("node") is None:
        pytest.skip("Node.js absent")
    code = _function("parsePartners") + """
    console.log(JSON.stringify(parsePartners("Sunu Assurances — Plateau, Abidjan\\n\\n  NSIA - Cocody ;  Riviera\\nAllianz\\n")));"""
    result = json.loads(subprocess.run(["node", "-e", code], capture_output=True, text=True, check=True).stdout)
    assert result == [{"name": "Sunu Assurances", "address": "Plateau, Abidjan"},
                      {"name": "NSIA", "address": "Cocody, Riviera"}, {"name": "Allianz", "address": None}]
