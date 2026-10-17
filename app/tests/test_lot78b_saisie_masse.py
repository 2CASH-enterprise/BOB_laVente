"""
Lot 78 B — saisie en masse : modèle Excel par secteur et par catégorie, import Excel/CSV (en-têtes français, colonnes
de fiche, catégorie Bob), export réimportable, recherche dans les fiches, champs du catalogue Meta.
"""
import io

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.core.security import hash_password
from app.models.product import Product
from app.models.tenant import Tenant, TenantPlan
from app.models.user import Role, User
from app.repositories.product_repository import ProductRepository
from app.services import product_sheets
from app.services.business_type import CAR_DEALERSHIP, INSURANCE_BROKER, ONLINE_STORE, REAL_ESTATE
from app.services.catalog_import import _clean_price, table_to_import_csv


async def _tenant(db, email, business_type=ONLINE_STORE, currency="EUR", default_category=None):
    tenant = Tenant(name="Boutique", country="FR", currency=currency, email=email, plan=TenantPlan.PRO,
                    business_type=business_type, default_bob_category=default_category)
    db.add(tenant)
    await db.flush()
    db.add(User(tenant_id=tenant.id, email=email, hashed_password=hash_password("x"), full_name="U", role=Role.OWNER))
    await db.commit()
    return tenant


async def _auth(client, email):
    token = (await client.post("/api/v1/auth/login", data={"username": email, "password": "x"})).json()["access_token"]
    return {"Authorization": f"Bearer {token}"}


def _xlsx(rows) -> bytes:
    book = Workbook()
    for row in rows:
        book.active.append(row)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _sheet_rows(raw: bytes, title=None):
    book = load_workbook(io.BytesIO(raw))
    return [list(r) for r in (book[title] if title else book.active).iter_rows(values_only=True)], book.sheetnames


async def _import(client, headers, filename, raw):
    return await client.post("/api/v1/catalog/import-csv", headers=headers, files={"file": (filename, raw, "x")})


async def _products(db, tenant):
    tenant_id = tenant.id
    db.expire_all()
    rows = (await db.execute(select(Product).where(Product.tenant_id == tenant_id))).scalars().all()
    return {p.sku: p for p in rows}


# --- Conversion du tableau ---------------------------------------------------------------------------------------

def test_french_headers_become_canonical_and_prices_are_cleaned():
    text = table_to_import_csv(["Référence", "Nom", "Prix", "Quantité", "Matière", "Catégorie Bob"],
                               [["R1", "Robe", "1 250,50 €", "3", "coton", "FASHION"], ["", "", "", "", "", ""]],
                               default_currency="EUR")
    lines = text.strip().splitlines()
    assert lines[0] == "SKU,NAME,PRICE,STOCK,Matière,Catégorie Bob,CURRENCY"
    assert lines[1] == "R1,Robe,1250.50,3,coton,FASHION,EUR" and len(lines) == 2  # ligne vide ignorée
    assert _clean_price("19,90") == "19.90" and _clean_price("1,250.50") == "1250.50" and _clean_price("25000") == "25000"
    assert _clean_price("25 000 FCFA") == "25000"


def test_sheet_columns_kept_only_for_a_store_and_property_type_is_not_the_free_category():
    store = table_to_import_csv(["SKU", "Matière", "Quantity", "Contenance / poids"], [["A", "x", "2", "1 L"]])
    assert store.splitlines()[0] == "SKU,Matière,STOCK,Contenance / poids"  # « Quantity » est le stock
    estate = table_to_import_csv(["Référence", "Type", "Ville"], [["B", "Maison", "Lyon"]], keep_property=True)
    assert estate.splitlines()[0] == "SKU,Type,Ville"
    shop = table_to_import_csv(["SKU", "Type"], [["C", "Robes"]])
    assert shop.splitlines()[0] == "SKU,CATEGORY"
    dealer = table_to_import_csv(["SKU", "Marque", "Couleur"], [["D", "Peugeot", "Gris"]], keep_vehicle=True)
    assert dealer.splitlines()[0] == "SKU,Marque,Couleur"


def test_flatten_sheet_and_variants():
    text = product_sheets.flatten({"material": "coton", "brand": "Nike", "compare_at_price": 30},
                                  [{"options": {"Taille": "M", "Couleur": "Rouge"}}, {"options": {"Taille": "M"}}])
    assert text == "coton | Nike | 30 | M | Rouge"
    assert product_sheets.flatten(None, None) is None
    assert len(product_sheets.flatten({"features": "x" * 5000})) == 4000


def test_category_from_code_or_label():
    assert product_sheets.category_from_value("fashion") == "FASHION"
    assert product_sheets.category_from_value(product_sheets.label("BEAUTY")) == "BEAUTY"
    assert product_sheets.category_from_value("Voitures") is None and product_sheets.category_from_value("") is None


# --- Import ----------------------------------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_excel_import_fills_the_sheet_and_category(client, db_session, unique_email):
    tenant = await _tenant(db_session, unique_email)
    headers = await _auth(client, unique_email)
    raw = _xlsx([["Référence", "Nom", "Prix", "Stock", "CATEGORIE_BOB", "Matière", "Tailles ou pointures", "Pour"],
                 ["R1", "Robe d'été", 39.9, 5, "FASHION", "lin", "S, M, L", "Femme"],
                 ["R2", "Crème", "12,50", 2, "", "", "", ""],
                 ["R3", "Montre", 99, 1, "Voitures", "", "", ""]])
    report = (await _import(client, headers, "catalogue.xlsx", raw)).json()
    assert report["imported"] == 2 and report["failed"] == 1
    assert "catégorie Bob inconnue « Voitures »" in report["errors"][0]
    products = await _products(db_session, tenant)
    robe = products["R1"]
    assert robe.bob_category == "FASHION" and robe.currency == "EUR" and float(robe.price) == 39.9
    assert robe.sheet == {"material": "lin", "sizes": "S, M, L", "gender": "Femme"}
    assert "lin" in robe.sheet_text
    assert products["R2"].bob_category == "BEAUTY" and products["R2"].sheet is None  # devinée d'après le nom


@pytest.mark.asyncio
async def test_reimport_replaces_only_the_columns_present(client, db_session, unique_email):
    tenant = await _tenant(db_session, unique_email)
    db_session.add(Product(tenant_id=tenant.id, sku="R1", name="Article", price=30, currency="EUR", stock_quantity=1,
                           bob_category="FASHION", sheet={"material": "lin", "colors": "Bleu", "care": "30°"}))
    await db_session.commit()
    headers = await _auth(client, unique_email)
    csv = "SKU;NAME;PRICE;CURRENCY;STOCK;Matière;Couleurs\nR1;Article;30;EUR;1;soie;\n"
    report = (await _import(client, headers, "c.csv", csv.encode("cp1252"))).json()
    assert report["updated"] == 1, report
    robe = (await _products(db_session, tenant))["R1"]
    # La catégorie du commerçant reste (pas de colonne CATEGORIE_BOB) ; Couleurs vide efface ; Entretien absent reste.
    assert robe.bob_category == "FASHION"
    assert robe.sheet == {"material": "soie", "care": "30°"}


@pytest.mark.asyncio
async def test_reimport_with_category_column_changes_the_category(client, db_session, unique_email):
    tenant = await _tenant(db_session, unique_email)
    db_session.add(Product(tenant_id=tenant.id, sku="R1", name="Robe", price=30, currency="EUR", stock_quantity=1,
                           bob_category="HOME", sheet={"care": "30°"}))
    await db_session.commit()
    headers = await _auth(client, unique_email)
    csv = "SKU,NAME,PRICE,CURRENCY,CATEGORIE_BOB,Matière\nR1,Robe,30,EUR,FASHION,soie\n"
    assert (await _import(client, headers, "c.csv", csv.encode())).json()["updated"] == 1
    robe = (await _products(db_session, tenant))["R1"]
    assert robe.bob_category == "FASHION" and robe.sheet == {"material": "soie", "care": "30°"}


@pytest.mark.asyncio
async def test_default_category_used_when_the_name_says_nothing(client, db_session, unique_email):
    tenant = await _tenant(db_session, unique_email, default_category="FOOD")
    headers = await _auth(client, unique_email)
    csv = "SKU,NAME,PRICE,CURRENCY\nZ1,Article mystère,5,EUR\n"
    assert (await _import(client, headers, "c.csv", csv.encode())).json()["imported"] == 1
    assert (await _products(db_session, tenant))["Z1"].bob_category == "FOOD"


@pytest.mark.asyncio
async def test_sheet_columns_ignored_outside_commerce(client, db_session, unique_email):
    tenant = await _tenant(db_session, unique_email, business_type=CAR_DEALERSHIP)
    headers = await _auth(client, unique_email)
    csv = "SKU,NAME,PRICE,CURRENCY,CATEGORIE_BOB,Matière,Marque\nV1,308,15000,EUR,FASHION,cuir,Peugeot\n"
    assert (await _import(client, headers, "c.csv", csv.encode())).json()["imported"] == 1
    car = (await _products(db_session, tenant))["V1"]
    assert car.bob_category is None and car.sheet is None and car.vehicle["brand"] == "Peugeot"


@pytest.mark.asyncio
async def test_unsupported_or_broken_files(client, db_session, unique_email):
    await _tenant(db_session, unique_email)
    headers = await _auth(client, unique_email)
    assert (await _import(client, headers, "c.txt", b"x")).status_code == 400
    assert (await _import(client, headers, "c.xls", b"x")).status_code == 400
    broken = await _import(client, headers, "c.xlsx", b"x")
    assert broken.status_code == 422 and "illisible" in broken.json()["detail"]
    assert (await _import(client, headers, "c.csv", b"")).status_code == 400


# --- Modèle et export --------------------------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_store_template_per_category(client, db_session, unique_email):
    await _tenant(db_session, unique_email)
    headers = await _auth(client, unique_email)
    response = await client.get("/api/v1/products/template/xlsx", params={"category": "FASHION"}, headers=headers)
    assert response.status_code == 200 and "modele-catalogue-bob-fashion.xlsx" in response.headers["content-disposition"]
    rows, sheets = _sheet_rows(response.content)
    header = rows[0]
    assert sheets == ["Catalogue", "Aide"] and len(rows) == 1
    assert header[:9] == product_sheets.BASE_COLUMNS + ["CATEGORIE_BOB"]
    assert "Tailles ou pointures" in header and "Matière" in header and "Ingrédients (INCI)" not in header
    help_rows, _ = _sheet_rows(response.content, "Aide")
    assert any("FASHION" in (r[0] or "") for r in help_rows)
    every = _sheet_rows((await client.get("/api/v1/products/template/xlsx", headers=headers)).content)[0][0]
    assert "Ingrédients (INCI)" in every and len(every) == 9 + len(product_sheets.field_columns())
    bad = await client.get("/api/v1/products/template/xlsx", params={"category": "CARS"}, headers=headers)
    assert bad.status_code == 400


@pytest.mark.asyncio
@pytest.mark.parametrize("business_type, expected, absent", [
    (REAL_ESTATE, ["TYPE", "TRANSACTION", "VILLE", "DPE", "REFERENCE MANDAT"], ["CATEGORIE_BOB"]),
    (CAR_DEALERSHIP, ["MARQUE", "KILOMETRAGE"], ["CATEGORIE_BOB", "TYPE"]),
    (INSURANCE_BROKER, ["SKU", "NAME"], ["PRICE", "CURRENCY", "STOCK", "CATEGORIE_BOB"]),
])
async def test_template_follows_the_sector(client, db_session, unique_email, business_type, expected, absent):
    await _tenant(db_session, unique_email, business_type=business_type)
    headers = await _auth(client, unique_email)
    header = _sheet_rows((await client.get("/api/v1/products/template/xlsx", headers=headers)).content)[0][0]
    assert all(h in header for h in expected) and not any(h in header for h in absent)


@pytest.mark.asyncio
async def test_store_export_reimports_unchanged(client, db_session, unique_email):
    tenant = await _tenant(db_session, unique_email)
    db_session.add_all([
        Product(tenant_id=tenant.id, sku="R1", name="Robe", price=39.9, currency="EUR", stock_quantity=2,
                bob_category="FASHION", sheet={"material": "lin", "sizes": "S, M"}),
        Product(tenant_id=tenant.id, sku="C1", name="Crème", price=12, currency="EUR", stock_quantity=0, active=False,
                bob_category="BEAUTY", sheet={"skin_hair_type": "sèche"}),
    ])
    await db_session.commit()
    headers = await _auth(client, unique_email)
    exported = (await client.get("/api/v1/products/export/xlsx", headers=headers)).content
    rows, _ = _sheet_rows(exported)
    header = rows[0]
    assert header[:9] == product_sheets.BASE_COLUMNS + ["CATEGORIE_BOB"]
    by_sku = {r[0]: dict(zip(header, r)) for r in rows[1:]}
    assert set(by_sku) == {"R1", "C1"}  # les produits désactivés aussi
    assert by_sku["R1"]["Matière"] == "lin" and by_sku["R1"]["CATEGORIE_BOB"] == product_sheets.label("FASHION")
    assert by_sku["C1"]["ACTIVE"] == "non" and "Ingrédients (INCI)" in header and "Coupe" in header
    report = (await _import(client, headers, "catalogue-bob.xlsx", exported)).json()
    assert report["updated"] == 2 and report["failed"] == 0, report
    products = await _products(db_session, tenant)
    assert products["R1"].sheet == {"material": "lin", "sizes": "S, M"} and products["R1"].bob_category == "FASHION"
    assert products["C1"].active is False and products["C1"].sheet == {"skin_hair_type": "sèche"}
    assert float(products["R1"].price) == 39.9


@pytest.mark.asyncio
async def test_real_estate_export_reimports_unchanged(client, db_session, unique_email):
    tenant = await _tenant(db_session, unique_email, business_type=REAL_ESTATE)
    sheet = {"transaction": "Vente", "city": "Lyon", "surface": 68, "rooms": 3, "dpe": "C", "furnished": "Vide"}
    db_session.add(Product(tenant_id=tenant.id, sku="M1", name="T3 lumineux", price=250000, currency="EUR",
                           stock_quantity=1, bob_category="APPARTEMENT", sheet=sheet))
    await db_session.commit()
    headers = await _auth(client, unique_email)
    exported = (await client.get("/api/v1/products/export/xlsx", headers=headers)).content
    rows, _ = _sheet_rows(exported)
    line = dict(zip(rows[0], rows[1]))
    assert line["TYPE"] == "Appartement" and line["VILLE"] == "Lyon" and line["SURFACE"] == 68
    assert (await _import(client, headers, "biens.xlsx", exported)).json()["updated"] == 1
    product = (await _products(db_session, tenant))["M1"]
    assert product.bob_category == "APPARTEMENT" and product.sheet == sheet


# --- Recherche dans les fiches -------------------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_search_finds_sheet_values_word_by_word(db_session, unique_email):
    tenant = await _tenant(db_session, unique_email)
    db_session.add_all([
        Product(tenant_id=tenant.id, sku="R1", name="Robe longue", price=40, currency="EUR", stock_quantity=1,
                bob_category="FASHION", sheet={"material": "coton bio"},
                variants=[{"options": {"Taille": "XL"}, "price": 40}]),
        Product(tenant_id=tenant.id, sku="R2", name="Robe courte", price=40, currency="EUR", stock_quantity=1,
                bob_category="FASHION", sheet={"material": "soie"}),
    ])
    await db_session.commit()
    repo = ProductRepository(db_session)
    tenant_id = tenant.id

    async def skus(query):
        return sorted(p.sku for p in await repo.search(tenant_id, query))

    assert await skus("coton") == ["R1"]
    assert await skus("robe coton") == ["R1"]  # chaque mot dans un champ différent
    assert await skus("robe") == ["R1", "R2"]
    assert await skus("xl") == ["R1"]  # valeur d'une variante
    assert await skus("robe laine") == []
    product = (await _products(db_session, tenant))["R2"]
    product.sheet = {"material": "laine"}
    await db_session.commit()
    assert await skus("laine") == ["R2"]  # texte à plat recalculé à chaque modification


# --- Catalogue Meta -----------------------------------------------------------------------------------------------

def test_meta_fields_map_to_the_sheet():
    from app.integrations.ecommerce.meta_catalog_client import PRODUCT_FIELDS, map_meta_catalog_product

    for field in ("brand", "color", "size", "material", "gender", "gtin"):
        assert field in PRODUCT_FIELDS.split(",")
    mapped = map_meta_catalog_product({"id": "1", "name": "Robe", "price": "39,90 €", "currency": "EUR",
                                       "availability": "in stock", "brand": "Zeta", "color": "Rouge", "size": "M",
                                       "material": "lin", "gender": "female", "gtin": "3001", "product_type": "Robes"})
    assert mapped["sheet"] == {"brand": "Zeta", "colors": "Rouge", "sizes": "M", "material": "lin", "gender": "Femme",
                               "ean": "3001"}
    assert mapped["bob_category"] == "FASHION"
    assert map_meta_catalog_product({"id": "2", "name": "X", "gender": "other"})["sheet"] == {}


class _FakeMeta:
    def __init__(self, products):
        self.products = products

    async def fetch_products(self):
        return self.products


@pytest.mark.asyncio
async def test_meta_sync_fills_the_sheet_without_overwriting_the_merchant(db_session, unique_email):
    from app.models.ecommerce_connection import EcommerceConnection
    from app.services.meta_catalog_sync import sync_meta_catalog

    tenant = await _tenant(db_session, unique_email)
    connection = EcommerceConnection(tenant_id=tenant.id, platform="META_CATALOG", shop_domain="123", access_token="t")
    db_session.add(connection)
    await db_session.commit()
    raw = {"id": "9", "retailer_id": "M9", "name": "Robe", "price": "39,90 €", "currency": "EUR",
           "availability": "in stock", "color": "Rouge", "material": "lin"}
    tenant_id = tenant.id
    await sync_meta_catalog(db_session, tenant_id, connection, client=_FakeMeta([raw]))
    product = (await _products(db_session, tenant))["M9"]
    assert product.bob_category == "FASHION" and product.sheet == {"colors": "Rouge", "material": "lin"}
    product.bob_category, product.sheet = "HOME", {"material": "lin", "care": "à la main"}
    await db_session.commit()
    await sync_meta_catalog(db_session, tenant_id, connection, client=_FakeMeta([{**raw, "material": "coton"}]))
    product = (await _products(db_session, tenant))["M9"]
    assert product.bob_category == "HOME"  # choisie par le commerçant : jamais écrasée
    assert product.sheet == {"material": "coton", "care": "à la main"}  # « colors » n'est pas un champ de HOME


@pytest.mark.asyncio
async def test_meta_sync_leaves_other_sectors_alone(db_session, unique_email):
    from app.models.ecommerce_connection import EcommerceConnection
    from app.services.meta_catalog_sync import sync_meta_catalog

    tenant = await _tenant(db_session, unique_email, business_type=CAR_DEALERSHIP)
    connection = EcommerceConnection(tenant_id=tenant.id, platform="META_CATALOG", shop_domain="123", access_token="t")
    db_session.add(connection)
    await db_session.commit()
    await sync_meta_catalog(db_session, tenant.id, connection, client=_FakeMeta([
        {"id": "9", "retailer_id": "M9", "name": "Robe", "price": "39,90 €", "currency": "EUR",
         "availability": "in stock", "color": "Rouge"}]))
    product = (await _products(db_session, tenant))["M9"]
    assert product.bob_category is None and product.sheet is None


# --- Tableau de bord --------------------------------------------------------------------------------------------

def test_dashboard_accepts_excel_and_offers_the_template():
    html = open("app/static/dashboard/index.html", encoding="utf-8").read()
    assert 'accept=".xlsx,.csv' in html and "Importer un catalogue Excel ou CSV" in html
    assert 'onclick="downloadTemplate()"' in html and 'id="template-category"' in html
    assert "/api/v1/products/template/xlsx" in html and "\\.(csv|xlsx)$" in html
    importer = html[html.index("async function importCsv() {"):html.index("async function loadCompanyProfile()")]
    assert "result.errors" in importer
